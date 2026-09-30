"""Render the Chinese demonstration source to an editable Word guide (no case images)."""
import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def build(source, output):
    document = Document()
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (("Normal", 10.5), ("Title", 23), ("Heading 1", 17),
                       ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.18
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Cardiac-Agent | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "---":
            document.add_page_break()
        elif line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
            cells = [r for r in cells if not all(set(c) <= {"-", ":", " "} for c in r)]
            table = document.add_table(rows=0, cols=len(cells[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = [5, 5.2, 6.8] if len(cells[0]) == 3 else [17 / len(cells[0])] * len(cells[0])
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for row_index, row in enumerate(cells):
                tr = table.add_row()
                no_split = OxmlElement("w:cantSplit")
                tr._tr.get_or_add_trPr().append(no_split)
                for cell, value, width in zip(tr.cells, row, widths):
                    cell.width = Cm(width)
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = row_index == 0
                    if row_index == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E9EEF0")
                        cell._tc.get_or_add_tcPr().append(shade)
                if row_index == 0:
                    tr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            document.add_paragraph().paragraph_format.space_after = Pt(0)
        elif line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        else:
            document.add_paragraph(line)
    document.core_properties.title = next(
        line[2:].strip() for line in lines if line.startswith("# "))
    document.core_properties.author = "Cardiac-Agent 项目组"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
